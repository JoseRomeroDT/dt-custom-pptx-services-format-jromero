#!/usr/bin/env python3
"""
inspect_deck.py — Lista slides, formas y textos de un .pptx (skill dt-custom-pptx-services-format-jromero).

Para cada slide muestra: número, id de patrón (nombre del slide), layout y cada forma con su nombre,
tipo, posición/tamaño en pulgadas y sus párrafos. Dentro de un párrafo, los runs se separan con "¦"
(cada run puede tener formato distinto: color, negrita…).

Uso:
  python scripts/inspect_deck.py deck.pptx                 # todo
  python scripts/inspect_deck.py deck.pptx --slide 5       # un slide
  python scripts/inspect_deck.py deck.pptx --placeholders  # solo textos [entre corchetes] (catálogo)
  python scripts/inspect_deck.py deck.pptx --json          # salida estructurada
"""
import argparse, json, re, sys
from pptx import Presentation
from pptx.util import Emu

PLACEHOLDER = re.compile(r"\[[^\[\]]{3,}\]")


def walk(shapes, depth=0):
    for sh in shapes:
        yield sh, depth
        if sh.shape_type == 6:  # grupo
            yield from walk(sh.shapes, depth + 1)


def paragraphs(sh):
    out = []
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            runs = [r.text for r in p.runs]
            if any(t.strip() for t in runs):
                out.append("¦".join(runs))
    if getattr(sh, "has_table", False) and sh.has_table:
        for row in sh.table.rows:
            for cell in row.cells:
                for p in cell.text_frame.paragraphs:
                    runs = [r.text for r in p.runs]
                    if any(t.strip() for t in runs):
                        out.append("¦".join(runs))
    return out


def inch(v):
    return round(Emu(v).inches, 2) if v is not None else None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--slide", type=int)
    ap.add_argument("--placeholders", action="store_true")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except AttributeError:
        pass
    prs = Presentation(a.pptx)
    data = []
    for n, s in enumerate(prs.slides, 1):
        if a.slide and n != a.slide:
            continue
        item = {"n": n, "id": s._element.cSld.get("name", ""), "layout": s.slide_layout.name, "shapes": []}
        for sh, depth in walk(s.shapes):
            paras = paragraphs(sh)
            if a.placeholders:
                paras = [p for p in paras if PLACEHOLDER.search(p.replace("¦", ""))]
                if not paras and not (sh.shape_type == 13 and sh.name.startswith("diagram")):
                    continue
            item["shapes"].append({"name": sh.name, "type": str(sh.shape_type).split(".")[-1].split(" ")[0],
                                   "depth": depth, "x": inch(sh.left), "y": inch(sh.top),
                                   "w": inch(sh.width), "h": inch(sh.height), "text": paras})
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
            item["notes"] = s.notes_slide.notes_text_frame.text.strip()
        data.append(item)
    if a.json:
        print(json.dumps(data, ensure_ascii=False, indent=2))
        return
    for it in data:
        print(f"\n=== Slide {it['n']} · id={it['id'] or '-'} · layout={it['layout']!r}")
        for sh in it["shapes"]:
            pad = "  " * (sh["depth"] + 1)
            geo = f"({sh['x']},{sh['y']} {sh['w']}x{sh['h']}in)"
            print(f"{pad}- {sh['name']} [{sh['type']}] {geo}")
            for p in sh["text"]:
                print(f"{pad}    · {p}")
        if it.get("notes") and not a.placeholders:
            print(f"  notes: {it['notes'][:160]}{'…' if len(it['notes']) > 160 else ''}")


if __name__ == "__main__":
    main()
