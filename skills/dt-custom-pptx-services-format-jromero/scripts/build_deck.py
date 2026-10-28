#!/usr/bin/env python3
"""
build_deck.py — Arma una presentación con el template definitivo de Dynatrace Services
(skill dt-custom-pptx-services-format-jromero).

Toma un spec JSON con la lista de slides; cada una indica el patrón del template (su id) y lo que
se reemplaza. El script:
  1. Copia cada patrón en el orden pedido (se puede repetir un patrón) y borra los que no se usan.
  2. Reemplaza los textos [placeholder] conservando el formato del template (fuente, color, tamaño).
  3. Cambia las imágenes (diagramas) ajustándolas al marco sin deformarlas.
  4. Quita formas opcionales, pone las notas del presentador y limpia las notas del template.
  5. Revisa el resultado: placeholders sin reemplazar, claves del spec que no se usaron y
     diagramas de ejemplo que siguen en el deck.

Spec (JSON):
{
  "output": "salida.pptx",                      # opcional si se pasa -o
  "template": "ruta/al/template.pptx",          # opcional; por defecto assets/template-dt-services.pptx
  "slides": [
    {"pattern": "cover",
     "text": {"[Presentation title]": "From customer issue to action plan",
              "[Subtitle]": "AI-assisted root cause analysis"},
     "notes": "Lo que el presentador dice en esta slide."},
    {"pattern": "key-messages",
     "text": {"[Key message 1]": "...", "[Key message 2]": "...", "[Key message 3]": ""}},   # "" borra el párrafo
    {"pattern": "diagram-banner",
     "text": {"[Slide title: the takeaway in one sentence]": "A six-step journey"},
     "images": {"diagram": "diagrams/journey.png"},           # o {"path": "...", "alt": "texto alternativo"}
     "remove": ["Callout"]}                       # nombres de formas o textos [placeholder] a quitar
  ]
}

Reglas de reemplazo de "text":
- La clave es el texto del placeholder tal como aparece en el template (ver references/catalogo-patrones.md).
- Si la clave es un párrafo completo, el valor puede ser un texto o una lista de textos (un párrafo cada uno).
- Si la clave es un run dentro de un párrafo (p. ej. una etiqueta en color), se reemplaza solo ese run.
- Un valor "" elimina el párrafo (o deja vacío el run).

Uso:
  python scripts/build_deck.py spec.json -o salida.pptx [--template otro.pptx] [--strict]
  python scripts/build_deck.py --list-patterns            # ids de patrón disponibles en el template
"""
import argparse, copy, hashlib, json, re, sys
from pathlib import Path
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.text.text import _Paragraph
from pptx.util import Emu

SKILL = Path(__file__).resolve().parent.parent
DEFAULT_TEMPLATE = SKILL / "assets" / "template-dt-services.pptx"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
PLACEHOLDER = re.compile(r"\[[^\[\]]{3,}\]")
SKIP_RELS = {RT.SLIDE_LAYOUT, RT.NOTES_SLIDE}


# ───────────────────────────── copia y borrado de slides ─────────────────────────────
def pattern_id(slide):
    return slide._element.cSld.get("name", "")


def find_layout(dst_prs, src_layout, src_prs):
    """Layout equivalente en el deck destino: mismo índice de master y mismo nombre."""
    m_idx = list(src_prs.slide_masters).index(src_layout.slide_master)
    masters = list(dst_prs.slide_masters)
    if m_idx < len(masters):
        for lay in masters[m_idx].slide_layouts:
            if lay.name == src_layout.name:
                return lay
    for m in masters:
        for lay in m.slide_layouts:
            if lay.name == src_layout.name:
                return lay
    raise SystemExit(f"ERROR: el deck destino no tiene el layout {src_layout.name!r}.")


def clone_part(part, dst_package, cache):
    """Clona una parte simple (imagen, svg, tags) de otro paquete."""
    if id(part) in cache:
        return cache[id(part)]
    if len(part.rels):
        raise SystemExit(f"ERROR: no se puede copiar entre archivos la parte {part.partname} (tiene relaciones).")
    name = str(part.partname)
    tmpl = re.sub(r"\d+(\.\w+)$", r"%d\1", name)
    new = Part.load(dst_package.next_partname(tmpl), part.content_type, dst_package, part.blob)
    cache[id(part)] = new
    return new


def copy_slide(src_slide, src_prs, dst_prs, cache=None):
    """Copia un slide (formas, fondo y relaciones) al final de dst_prs. Sirve en el mismo archivo o entre archivos."""
    cache = {} if cache is None else cache
    same_pkg = src_prs is dst_prs
    layout = src_slide.slide_layout if same_pkg else find_layout(dst_prs, src_slide.slide_layout, src_prs)
    new = dst_prs.slides.add_slide(layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    src_csld, new_csld = src_slide._element.cSld, new._element.cSld
    bg = src_csld.find(qn("p:bg"))
    if bg is not None:
        new_csld.insert(0, copy.deepcopy(bg))
    rid_map = {}
    for rId, rel in src_slide.part.rels.items():
        if rel.reltype in SKIP_RELS:
            continue
        if rel.is_external:
            rid_map[rId] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            target = rel.target_part if same_pkg else clone_part(rel.target_part, dst_prs.part.package, cache)
            rid_map[rId] = new.part.relate_to(target, rel.reltype)
    tree = new.shapes._spTree
    for el in src_slide.shapes._spTree.iterchildren():
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        tree.append(copy.deepcopy(el))
    for el in new._element.iter():
        for attr, val in list(el.attrib.items()):
            if attr.startswith(R_NS) and val in rid_map:
                el.set(attr, rid_map[val])
    for tag in ("p:transition", "p:timing"):
        src_el = src_slide._element.find(qn(tag))
        if src_el is not None:
            new._element.append(copy.deepcopy(src_el))
    if src_csld.get("name"):
        new_csld.set("name", src_csld.get("name"))
    return new


def delete_slide(prs, slide):
    lst = prs.slides._sldIdLst
    for sld in list(lst):
        if prs.part.related_part(sld.rId) is slide.part:
            rId = sld.rId
            lst.remove(sld)
            prs.part.drop_rel(rId)
            return
    raise SystemExit("ERROR: slide no encontrado para borrar.")


# ───────────────────────────── texto ─────────────────────────────
def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from walk(sh.shapes)


def text_frames(slide):
    for sh in walk(slide.shapes):
        if sh.has_text_frame:
            yield sh.text_frame
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


def set_paragraph(p, text):
    runs = p.runs
    if not runs:
        p.add_run().text = text
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def put(p, key, value, whole):
    """Escribe value en lugar de key dentro del párrafo, conservando el formato de su run."""
    for r in p.runs:
        if key in r.text:
            r.text = r.text.replace(key, value)
            return
    if whole:  # la clave abarca varios runs: queda el formato del primero
        set_paragraph(p, value)


def count_paragraphs(tf):
    return sum(1 for p in tf.paragraphs if "".join(r.text for r in p.runs).strip())


def grow_boxes(slide, before):
    """Si una lista sumó párrafos, la caja crece en proporción: PowerPoint no reajusta el tamaño guardado."""
    for sh in walk(slide.shapes):
        if not sh.has_text_frame or sh._element not in before:
            continue
        old, new = before[sh._element], count_paragraphs(sh.text_frame)
        if new > old > 0:  # también con autoajuste: PowerPoint no lo recalcula al abrir el archivo
            sh.height = int(sh.height * new / old)


def replace_text(slide, mapping, used):
    keys = sorted(mapping, key=len, reverse=True)
    before = {sh._element: count_paragraphs(sh.text_frame) for sh in walk(slide.shapes) if sh.has_text_frame}
    for tf in text_frames(slide):
        for p in list(tf.paragraphs):
            for key in keys:
                full = "".join(r.text for r in p.runs)
                if key not in full:
                    continue
                used.add(key)
                val = mapping[key]
                items = [str(v) for v in val] if isinstance(val, list) else [str(val)]
                whole = full.strip() == key
                if not items or items == [""]:
                    if whole or isinstance(val, list) or not full.replace(key, "").strip(" 🎯•-–—:"):
                        if len(tf.paragraphs) > 1:
                            p._p.getparent().remove(p._p)
                        else:
                            set_paragraph(p, "")
                        break
                    put(p, key, "", whole)
                    continue
                anchor = p._p
                for extra in items[1:]:  # un párrafo más por cada elemento de la lista
                    clone = copy.deepcopy(p._p)
                    anchor.addnext(clone)
                    anchor = clone
                    put(_Paragraph(clone, tf), key, extra, whole)
                put(p, key, items[0], whole)
    grow_boxes(slide, before)


# ───────────────────────────── imágenes y formas ─────────────────────────────
def replace_image(slide, shape_name, path, alt=None):
    pics = [sh for sh in walk(slide.shapes) if sh.name == shape_name and sh.shape_type == 13]
    if not pics:
        raise SystemExit(f"ERROR: el patrón {pattern_id(slide)!r} no tiene una imagen llamada {shape_name!r}.")
    pic = pics[0]
    image_part, rId = slide.part.get_or_add_image_part(str(path))
    blip = pic._element.blipFill.find(qn("a:blip"))
    old_rId = blip.get(qn("r:embed"))
    blip.set(qn("r:embed"), rId)
    if old_rId and old_rId != rId:
        slide.part.drop_rel(old_rId)  # la imagen de ejemplo no viaja en el deck si nada más la usa
    pic._element.nvPicPr.cNvPr.set("descr", alt or Path(path).stem.replace("-", " ").replace("_", " "))
    src = pic._element.blipFill.find(qn("a:srcRect"))
    if src is not None:
        src.getparent().remove(src)
    iw, ih = image_part.image.size
    fx, fy, fw, fh = pic.left, pic.top, pic.width, pic.height
    scale = min(fw / iw, fh / ih)
    w, h = int(iw * scale), int(ih * scale)
    pic.left, pic.top, pic.width, pic.height = fx + (fw - w) // 2, fy + (fh - h) // 2, w, h


def remove_shapes(slide, targets):
    done = set()
    for t in targets:
        hit = False
        for sh in list(walk(slide.shapes)):
            txt = sh.text_frame.text if sh.has_text_frame else ""
            if sh.name == t or (PLACEHOLDER.fullmatch(t) and t in txt):
                el = sh._element
                if el.getparent() is not None:
                    el.getparent().remove(el)
                hit = True
        if hit:
            done.add(t)
    return set(targets) - done


def set_notes(slide, text):
    tf = slide.notes_slide.notes_text_frame
    tf.text = text or ""


# ───────────────────────────── revisión ─────────────────────────────
def leftovers(slide):
    out = []
    for tf in text_frames(slide):
        for p in tf.paragraphs:
            out += PLACEHOLDER.findall("".join(r.text for r in p.runs))
    return out


def image_hashes(slide):
    hs = {}
    for sh in walk(slide.shapes):
        if sh.shape_type == 13 and sh.name.startswith("diagram"):
            hs[sh.name] = hashlib.sha1(sh.image.blob).hexdigest()
    return hs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("spec", nargs="?")
    ap.add_argument("-o", "--output")
    ap.add_argument("--template")
    ap.add_argument("--strict", action="store_true", help="las advertencias también bloquean")
    ap.add_argument("--list-patterns", action="store_true")
    a = ap.parse_args()
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except AttributeError:
        pass

    spec = {}
    if a.spec:
        spec_path = Path(a.spec).resolve()
        spec = json.loads(spec_path.read_text(encoding="utf-8"))
        base = spec_path.parent
    else:
        base = Path.cwd()
    tpl = Path(a.template or spec.get("template") or DEFAULT_TEMPLATE)
    if not tpl.is_absolute() and not tpl.exists():
        tpl = base / tpl
    prs = Presentation(str(tpl))
    patterns = {pattern_id(s): s for s in prs.slides if pattern_id(s)}
    if a.list_patterns:
        for i, (pid, s) in enumerate(patterns.items(), 1):
            print(f"{i:2d}. {pid}")
        return
    if not a.spec:
        ap.error("falta el spec JSON")
    out = Path(a.output or spec.get("output") or "deck.pptx")
    if not out.is_absolute() and not a.output:
        out = base / out

    errors, warnings = [], []
    template_slides = list(prs.slides)
    tpl_hashes = {pid: image_hashes(s) for pid, s in patterns.items()}
    built = []
    for i, item in enumerate(spec.get("slides", []), 1):
        pid = item.get("pattern")
        if pid not in patterns:
            errors.append(f"Slide {i}: el patrón {pid!r} no existe. Usa --list-patterns.")
            continue
        new = copy_slide(patterns[pid], prs, prs)
        built.append((i, pid, new, item))
    if errors:
        print("ERRORES:\n - " + "\n - ".join(errors), file=sys.stderr)
        sys.exit(1)
    for s in template_slides:
        delete_slide(prs, s)

    for i, pid, slide, item in built:
        used = set()
        mapping = {k.strip(): v for k, v in item.get("text", {}).items()}
        missing = remove_shapes(slide, item.get("remove", []))
        for m in missing:
            warnings.append(f"Slide {i} ({pid}): no se encontró la forma a quitar {m!r}.")
        replace_text(slide, mapping, used)
        for k in mapping:
            if k not in used:
                warnings.append(f"Slide {i} ({pid}): la clave {k!r} no está en el patrón.")
        for name, val in item.get("images", {}).items():
            path, alt = (val.get("path"), val.get("alt")) if isinstance(val, dict) else (val, None)
            p = Path(path)
            p = p if p.is_absolute() else base / p
            if not p.exists():
                errors.append(f"Slide {i} ({pid}): no existe la imagen {p}.")
                continue
            replace_image(slide, name, p, alt)
        set_notes(slide, item.get("notes", ""))
        left = leftovers(slide)
        if left:
            warnings.append(f"Slide {i} ({pid}): placeholders sin reemplazar: {', '.join(sorted(set(left))[:6])}")
        for name, h in image_hashes(slide).items():
            if tpl_hashes.get(pid, {}).get(name) == h:
                warnings.append(f"Slide {i} ({pid}): la imagen {name!r} sigue siendo el diagrama de ejemplo del template.")
        if not item.get("notes"):
            warnings.append(f"Slide {i} ({pid}): sin notas del presentador.")

    if errors:
        print("ERRORES:\n - " + "\n - ".join(errors), file=sys.stderr)
        sys.exit(1)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    if warnings:
        print("ADVERTENCIAS:\n - " + "\n - ".join(warnings), file=sys.stderr)
    print(f"OK: {out} | slides: {len(built)} | patrones: {', '.join(pid for _, pid, _, _ in built)}")
    if a.strict and warnings:
        print("BLOQUEADO POR --strict: corrige las advertencias.", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
