#!/usr/bin/env python3
"""
check_deck.py — Control de calidad de una presentación hecha con el template de Dynatrace Services
(skill dt-custom-pptx-services-format-jromero). No modifica el archivo.

Errores (bloquean la entrega):
  - Placeholders sin reemplazar: textos [entre corchetes] del template.
  - Diagramas de ejemplo del template que siguen en el deck (imágenes 'diagram' sin cambiar).
Advertencias:
  - Slides sin notas del presentador.
  - Fuentes fuera de la familia DT Flow.
  - Títulos de más de 70 caracteres (deben caber en una línea).
  - Textos que, por cantidad de caracteres, probablemente no caben en su caja (revisar en el render).

Uso:
  python scripts/check_deck.py deck.pptx [--template assets/template-dt-services.pptx] [--strict]
"""
import argparse, hashlib, re, sys
from pathlib import Path
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

SKILL = Path(__file__).resolve().parent.parent
DEFAULT_TEMPLATE = SKILL / "assets" / "template-dt-services.pptx"
PLACEHOLDER = re.compile(r"\[[^\[\]]{3,}\]")
FONT_OK = re.compile(r"^(DT Flow.*|\+m[jn]-.*|Segoe UI Emoji|Segoe UI Symbol)$")


def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from walk(sh.shapes)


def diagram_hashes(prs):
    out = set()
    for s in prs.slides:
        for sh in walk(s.shapes):
            if sh.shape_type == 13 and sh.name.startswith("diagram"):
                out.add(hashlib.sha1(sh.image.blob).hexdigest())
    return out


def fits(sh):
    """Estimación gruesa, párrafo por párrafo: alto que necesita el texto frente al alto de la caja (en puntos)."""
    if not sh.has_text_frame or not sh.width or not sh.height:
        return True, 0, 0
    if not sh.text_frame.text.strip():
        return True, 0, 0
    if sh.text_frame._txBody.bodyPr.find(qn("a:spAutoFit")) is not None:
        return True, 0, 0  # la caja crece con el texto
    width_pt = max(1.0, (Emu(sh.width).inches - 0.2) * 72)
    box_pt = Emu(sh.height).inches * 72
    need_pt = 0.0
    for p in sh.text_frame.paragraphs:
        text = "".join(r.text for r in p.runs)
        sizes = [r.font.size.pt for r in p.runs if r.font.size]
        pt = max(sizes) if sizes else 14
        if not text.strip():
            continue
        per_line = max(1, int(width_pt / (pt * 0.47)))  # ancho medio de carácter de DT Flow
        need_pt += max(1, -(-len(text) // per_line)) * pt * 1.2
    return need_pt <= box_pt * 1.35, round(need_pt), round(box_pt)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--template", default=str(DEFAULT_TEMPLATE))
    ap.add_argument("--strict", action="store_true")
    a = ap.parse_args()
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except AttributeError:
        pass
    prs = Presentation(a.pptx)
    tpl_hashes = diagram_hashes(Presentation(a.template)) if Path(a.template).exists() else set()
    errors, warnings = [], []
    for n, s in enumerate(prs.slides, 1):
        pid = s._element.cSld.get("name", "") or "-"
        tag = f"Slide {n} ({pid})"
        left = set()
        for sh in walk(s.shapes):
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    left.update(PLACEHOLDER.findall("".join(r.text for r in p.runs)))
                    for r in p.runs:
                        f = r.font.name
                        if f and not FONT_OK.match(f):
                            warnings.append(f"{tag}: fuente '{f}' fuera de la familia DT Flow en '{r.text[:30]}'.")
                ok, need, box = fits(sh)
                if not ok:
                    snippet = sh.text_frame.text.strip().replace("\n", " / ")[:40]
                    warnings.append(f"{tag}: '{snippet}…' en '{sh.name}' podría no caber "
                                    f"(necesita ~{need} pt de alto, la caja tiene {box} pt). Revisa el render.")
            if sh.shape_type == 13 and sh.name.startswith("diagram"):
                if hashlib.sha1(sh.image.blob).hexdigest() in tpl_hashes:
                    errors.append(f"{tag}: la imagen '{sh.name}' sigue siendo el diagrama de ejemplo del template.")
            if sh.is_placeholder and sh.placeholder_format.type is not None and "TITLE" in str(sh.placeholder_format.type):
                t = sh.text_frame.text.strip()
                if len(t) > 70:
                    warnings.append(f"{tag}: título de {len(t)} caracteres (máx. 70, una línea): '{t[:50]}…'")
        if left:
            errors.append(f"{tag}: placeholders sin reemplazar: {', '.join(sorted(left)[:8])}")
        notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ""
        if not notes:
            warnings.append(f"{tag}: sin notas del presentador.")
        elif notes.startswith("PATRÓN "):
            errors.append(f"{tag}: tiene las notas de documentación del template; reemplázalas por las del presentador.")
    for w in warnings:
        print(f"AVISO  {w}")
    for e in errors:
        print(f"ERROR  {e}")
    print(f"RESULTADO: {len(prs.slides)} slides · {len(errors)} errores · {len(warnings)} avisos")
    if errors or (a.strict and warnings):
        sys.exit(1)


if __name__ == "__main__":
    main()
